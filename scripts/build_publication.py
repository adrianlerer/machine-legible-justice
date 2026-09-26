#!/usr/bin/env python3
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "manuscript.md"
OUT_DIR = ROOT / "release"
DOCX = OUT_DIR / "Machine-Legible-Justice-Lerer-2026.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_rich_text(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        else:
            paragraph.add_run(part)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)
    pf.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 14, 14, 6),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        st = styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Abstract Body" not in styles:
        abstract = styles.add_style("Abstract Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abstract = styles["Abstract Body"]
    abstract.font.name = "Times New Roman"
    abstract.font.size = Pt(10)
    abstract.font.color.rgb = RGBColor(0, 0, 0)
    abstract._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.line_spacing = 1.05
    abstract.paragraph_format.space_after = Pt(5)

    if "Reference" not in styles:
        ref = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        ref = styles["Reference"]
    ref.font.name = "Times New Roman"
    ref.font.size = Pt(9.5)
    ref.font.color.rgb = RGBColor(0, 0, 0)
    ref._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ref.paragraph_format.left_indent = Cm(0.6)
    ref.paragraph_format.first_line_indent = Cm(-0.6)
    ref.paragraph_format.space_after = Pt(4)


def add_table(doc, rows):
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, "D9D9D9")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value.strip())
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5)
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_rich_text(p, value.strip())
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.8)
    configure_styles(doc)
    add_page_number(sec.footer.paragraphs[0])

    abstract_mode = False
    references_mode = False
    table_rows = []
    front_matter_lines = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(s.strip() for s in paragraph_buffer).strip()
        paragraph_buffer = []
        if not text:
            return
        style = "Reference" if references_mode else ("Abstract Body" if abstract_mode else "Normal")
        p = doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not abstract_mode and not references_mode:
            p.paragraph_format.first_line_indent = Cm(0.6)
        add_rich_text(p, text)
        if abstract_mode:
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    for line in lines:
        if line.startswith("|"):
            flush_paragraph()
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
                continue
            table_rows.append(cells)
            continue
        flush_table()

        if not line.strip():
            flush_paragraph()
            continue

        if line.startswith("# "):
            flush_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
            r = p.add_run(line[2:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(20)
            front_matter_lines += 1
            continue
        if line.startswith("## "):
            flush_paragraph()
            text = line[3:].strip()
            if front_matter_lines == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(12)
                r = p.add_run(text)
                r.italic = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                front_matter_lines += 1
            else:
                abstract_mode = text == "Abstract"
                references_mode = text == "References"
                p = doc.add_paragraph(text, style="Heading 1")
                if abstract_mode:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        if line.startswith("### "):
            flush_paragraph()
            abstract_mode = False
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            continue

        if front_matter_lines == 2 and not line.startswith("Conceptual preprint"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            add_rich_text(p, line.strip())
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10.5)
            continue
        if line.startswith("Conceptual preprint"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(line.strip())
            r.italic = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(9)
            front_matter_lines = 3
            continue

        if line.startswith("**Keywords:**"):
            flush_paragraph()
            p = doc.add_paragraph(style="Abstract Body")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich_text(p, line.strip())
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
            abstract_mode = False
            continue

        paragraph_buffer.append(line)

    flush_paragraph()
    flush_table()

    props = doc.core_properties
    props.title = "Machine-Legible Justice"
    props.subject = "Adjudicative bias, synthetic credibility, and recursive lock-in in AI agents"
    props.author = "Ignacio Adrián Lerer"
    props.keywords = "adjudicative AI; procedural justice; LLM-as-a-judge; machine legibility"
    props.comments = "Conceptual preprint prepared for author review prior to repository deposit."
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
